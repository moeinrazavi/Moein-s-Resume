#!/usr/bin/env python3
"""Behavioral / methodology interview deck for the Charlotte Wang & Kate Landman round.

This round is about HOW I work, not a project tour. Each slide answers one of the
questions they shared, grounded in real Ford experience:

  1. General introduction
  2. How I measure success in my current role
  3. Backtesting recommendation models — the metrics I rely on
  4. Turning unstructured data into model features
  5. Deciding which features are worth the effort
  6. Hard conversations with VIP stakeholders (and keeping the relationship)
  7. Prioritizing feedback once it's collected
  8. Working style / team configuration / the role I take
  9. Agile vs waterfall (and helping a team make the move)

Same visual language as the simple portfolio so the personal brand is consistent.
"""

from __future__ import annotations

from pathlib import Path

import requests
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "presentation-assets"
OUTPUT = ROOT / "Moein-Razavi-Interview-Charlotte-Kate.pptx"

# ---------------------------------------------------------------------------
# Palette (matches the simple portfolio)
# ---------------------------------------------------------------------------
NAVY = RGBColor(20, 42, 78)
NAVY_DARK = RGBColor(13, 27, 51)
INK = RGBColor(28, 36, 54)
BODY = RGBColor(70, 82, 102)
MUTED = RGBColor(130, 140, 158)
LINE = RGBColor(220, 226, 236)
CREAM = RGBColor(250, 248, 244)
CORAL = RGBColor(232, 96, 80)
GOLD = RGBColor(232, 178, 70)
TEAL = RGBColor(52, 152, 162)
PLUM = RGBColor(150, 90, 158)
WHITE = RGBColor(255, 255, 255)
SOFT_WHITE = RGBColor(238, 244, 252)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

IMAGE_URLS = {
    "hero": "https://images.unsplash.com/photo-1454165804606-c3d57bc86b40?w=1920&q=80",
    "team": "https://images.unsplash.com/photo-1522071820081-009f0129c71c?w=1920&q=80",
}


def load_images() -> dict[str, Path]:
    ASSETS.mkdir(exist_ok=True)
    paths: dict[str, Path] = {}
    s = requests.Session()
    s.headers["User-Agent"] = "Moein-Interview-Deck/1.0"
    for key, url in IMAGE_URLS.items():
        dest = ASSETS / f"{key}.jpg"
        if not dest.exists():
            try:
                r = s.get(url, timeout=30)
                r.raise_for_status()
                dest.write_bytes(r.content)
            except Exception as exc:
                print(f"[warn] {key}: {exc}")
                continue
        paths[key] = dest
    return paths


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def fill(shape, color, transparency=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if transparency is not None:
        shape.fill.transparency = transparency


def rect(slide, l, t, w, h, color, transparency=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color, transparency)
    return s


def round_rect(slide, l, t, w, h, color, corner=0.05, transparency=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color, transparency)
    try:
        s.adjustments[0] = corner
    except Exception:
        pass
    return s


def text(slide, l, t, w, h, content, *, size=14, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=None, italic=False, font=None, wrap=True):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = content
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    if font:
        r.font.name = font
    return box


def bullets(slide, l, t, w, h, items, *, size=12, color=BODY,
            bullet=CORAL, space=7):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        r1 = p.add_run()
        r1.text = "•  "
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = bullet
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(size)
        r2.font.color.rgb = color


def image_overlay(slide, img: Path, transparency=0.55, color=NAVY_DARK):
    slide.shapes.add_picture(str(img), 0, 0, width=SLIDE_W, height=SLIDE_H)
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, color, transparency=transparency)


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


# ---------------------------------------------------------------------------
# Page chrome
# ---------------------------------------------------------------------------
def page(prs, *, title, subtitle=None, eyebrow=None, eyebrow_color=TEAL):
    s = blank(prs)
    bg(s, CREAM)
    rect(s, Inches(0.7), Inches(0.6), Inches(0.6), Inches(0.07), eyebrow_color)
    if eyebrow:
        text(s, Inches(0.7), Inches(0.72), Inches(11.9), Inches(0.3),
             eyebrow.upper(), size=11, bold=True, color=eyebrow_color)
        t_top = Inches(1.02)
    else:
        t_top = Inches(0.78)
    text(s, Inches(0.7), t_top, Inches(12), Inches(0.9),
         title, size=30, bold=True, color=NAVY)
    if subtitle:
        text(s, Inches(0.7), t_top + Inches(0.88), Inches(12), Inches(0.45),
             subtitle, size=14, color=BODY, italic=True)
    return s


def footer(slide, idx, total):
    rect(slide, Inches(0.7), Inches(7.05), Inches(11.93), Inches(0.012), LINE)
    text(slide, Inches(0.7), Inches(7.13), Inches(9), Inches(0.25),
         "Moein Razavi  ·  Interview with Charlotte Wang & Kate Landman",
         size=9, color=MUTED)
    text(slide, Inches(11.4), Inches(7.13), Inches(1.4), Inches(0.25),
         f"{idx} / {total}", size=9, color=MUTED, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# Reusable answer layouts
# ---------------------------------------------------------------------------
def value_card(s, x, y, w, h, header, items, accent,
               *, header_size=14, body_size=11.5):
    round_rect(s, x, y, w, h, WHITE, corner=0.05)
    rect(s, x, y, w, Inches(0.08), accent)
    text(s, x + Inches(0.25), y + Inches(0.24), w - Inches(0.5), Inches(0.55),
         header, size=header_size, bold=True, color=NAVY)
    bullets(s, x + Inches(0.25), y + Inches(0.92), w - Inches(0.5),
            h - Inches(1.05), items, size=body_size, color=BODY,
            bullet=accent, space=7)


def qa_columns(prs, *, idx, total, eyebrow, title, subtitle, columns, accent,
               bottom_note=None):
    """columns: list of (header, [bullets], accent_color)."""
    s = page(prs, eyebrow=eyebrow, eyebrow_color=accent, title=title,
             subtitle=subtitle)
    n = len(columns)
    gap = Inches(0.25)
    col_w = (Inches(11.93) - gap * (n - 1)) / n
    col_top = Inches(2.45)
    col_h = Inches(3.7) if bottom_note else Inches(4.35)
    for i, (header, items, caccent) in enumerate(columns):
        x = Inches(0.7) + i * (col_w + gap)
        value_card(s, x, col_top, col_w, col_h, header, items, caccent)
    if bottom_note:
        ny = col_top + col_h + Inches(0.18)
        round_rect(s, Inches(0.7), ny, Inches(11.93), Inches(0.62), NAVY,
                   corner=0.08)
        text(s, Inches(0.95), ny, Inches(11.4), Inches(0.62),
             bottom_note, size=12, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    footer(s, idx, total)
    return s


def qa_rows(prs, *, idx, total, eyebrow, title, subtitle, rows, accent):
    """rows: list of (lead, detail)."""
    s = page(prs, eyebrow=eyebrow, eyebrow_color=accent, title=title,
             subtitle=subtitle)
    top = Inches(2.5)
    row_h = Inches(0.6)
    gap = Inches(0.1)
    for i, (lead, detail) in enumerate(rows):
        y = top + i * (row_h + gap)
        round_rect(s, Inches(0.7), y, Inches(11.93), row_h, WHITE, corner=0.06)
        rect(s, Inches(0.7), y, Inches(0.6), row_h, accent)
        text(s, Inches(0.7), y, Inches(0.6), row_h, str(i + 1),
             size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        box = s.shapes.add_textbox(Inches(1.55), y, Inches(10.9), row_h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{lead}  "
        r1.font.size = Pt(12.5); r1.font.bold = True; r1.font.color.rgb = NAVY
        r2 = p.add_run()
        r2.text = f"— {detail}"
        r2.font.size = Pt(12); r2.font.color.rgb = BODY
    footer(s, idx, total)
    return s


# ---------------------------------------------------------------------------
# Bespoke slides
# ---------------------------------------------------------------------------
def slide_title(prs, images):
    s = blank(prs)
    if "hero" in images:
        image_overlay(s, images["hero"], transparency=0.58, color=NAVY_DARK)
    else:
        bg(s, NAVY_DARK)

    rect(s, 0, Inches(2.55), Inches(0.45), Inches(0.18), TEAL)
    text(s, Inches(0.85), Inches(2.4), Inches(11.5), Inches(0.4),
         "INTERVIEW  ·  PREPARED FOR CHARLOTTE WANG & KATE LANDMAN",
         size=13, bold=True, color=GOLD)
    text(s, Inches(0.85), Inches(2.95), Inches(11.5), Inches(1.3),
         "Hi, I'm Moein.", size=56, bold=True, color=WHITE)
    text(s, Inches(0.85), Inches(4.25), Inches(11.5), Inches(0.6),
         "Senior AI/ML Engineer — I ship production AI and work with the people who use it.",
         size=20, color=SOFT_WHITE)
    text(s, Inches(0.85), Inches(5.05), Inches(11.5), Inches(0.45),
         "Here's how I think about success, data, stakeholders, and the way I work.",
         size=15, italic=True, color=GOLD)
    text(s, Inches(0.85), Inches(6.5), Inches(11.5), Inches(0.35),
         "razavi.moein94@gmail.com  ·  linkedin.com/in/moein-razavi",
         size=11, color=SOFT_WHITE)
    notes(s, "Warm open. Thank Charlotte and Kate. This round is about how I work, "
             "so I'll keep examples concrete and let them steer.")


def slide_intro(prs, images, idx, total):
    s = page(prs, eyebrow="General introduction",
             title="A little about me", eyebrow_color=TEAL,
             subtitle="The short version, in plain English.")
    points = [
        "Senior AI/ML Engineer with a Ph.D. from Texas A&M (4.0 GPA).",
        "5+ years building production AI — currently at Ford Motor Company.",
        "I own systems end-to-end: from the business problem, to the model, to production and monitoring.",
        "I work at the seam between deep technical work and the business — I can go deep, and I can explain it.",
        "What I care about: measurable impact, earning trust, and actually shipping.",
    ]
    bullets(s, Inches(0.7), Inches(2.5), Inches(7.4), Inches(4.2),
            points, size=16, color=INK, bullet=TEAL, space=14)
    if "team" in images:
        s.shapes.add_picture(str(images["team"]),
                             Inches(8.5), Inches(2.5),
                             width=Inches(4.2), height=Inches(4.1))
    footer(s, idx, total)
    notes(s, "Keep it to ~60 seconds. Land the 'bridge between technical and business' point — "
             "it matters for this round.")


def slide_agenda(prs, idx, total):
    s = page(prs, eyebrow="What I'd love to cover",
             title="The questions you shared — my take on each",
             eyebrow_color=TEAL,
             subtitle="Stop me on any of them; happy to go deeper with real examples.")
    topics = [
        ("1", "Measuring success", "How I know my work is actually working", TEAL),
        ("2", "Backtesting recommendations", "The metrics I trust for rec models", NAVY),
        ("3", "Unstructured data → features", "Turning text and docs into model signals", CORAL),
        ("4", "Choosing what to build", "Deciding which features earn their keep", GOLD),
        ("5", "Hard conversations", "Disagreeing with senior stakeholders, keeping trust", PLUM),
        ("6", "Prioritizing feedback", "Turning a pile of feedback into a ranked plan", TEAL),
        ("7", "Working style", "How I show up on a team", NAVY),
        ("8", "Agile vs waterfall", "How I work — and helping the move to agile", CORAL),
    ]
    gap_x = Inches(0.25)
    gap_y = Inches(0.16)
    card_w = (Inches(11.93) - gap_x) / 2
    card_h = Inches(0.95)
    top = Inches(2.45)
    for i, (num, title, desc, accent) in enumerate(topics):
        col = i % 2
        row = i // 2
        x = Inches(0.7) + col * (card_w + gap_x)
        y = top + row * (card_h + gap_y)
        round_rect(s, x, y, card_w, card_h, WHITE, corner=0.07)
        rect(s, x, y, Inches(0.6), card_h, accent)
        text(s, x, y, Inches(0.6), card_h, num, size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(0.8), y + Inches(0.16), card_w - Inches(1.0), Inches(0.4),
             title, size=14, bold=True, color=NAVY)
        text(s, x + Inches(0.8), y + Inches(0.52), card_w - Inches(1.0), Inches(0.38),
             desc, size=11, color=BODY)
    footer(s, idx, total)


def slide_thanks(prs, images):
    s = blank(prs)
    if "hero" in images:
        image_overlay(s, images["hero"], transparency=0.62, color=NAVY_DARK)
    else:
        bg(s, NAVY_DARK)
    rect(s, Inches(6.0), Inches(2.0), Inches(1.3), Inches(0.18), TEAL)
    text(s, 0, Inches(2.4), SLIDE_W, Inches(1.2),
         "Thank you.", size=58, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, 0, Inches(3.7), SLIDE_W, Inches(0.5),
         "I'd love to hear how your team works — and answer anything.",
         size=18, color=GOLD, align=PP_ALIGN.CENTER)
    card_w = Inches(8.0)
    cx = (SLIDE_W - card_w) / 2
    round_rect(s, cx, Inches(4.7), card_w, Inches(1.85), NAVY_DARK, corner=0.04)
    text(s, cx, Inches(4.88), card_w, Inches(0.4), "MOEIN RAZAVI",
         size=12, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    text(s, cx, Inches(5.25), card_w, Inches(0.4),
         "razavi.moein94@gmail.com", size=14, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, cx, Inches(5.65), card_w, Inches(0.4),
         "(979) 676-7486", size=13, color=SOFT_WHITE, align=PP_ALIGN.CENTER)
    text(s, cx, Inches(6.05), card_w, Inches(0.4),
         "linkedin.com/in/moein-razavi  ·  github.com/moeinrazavi",
         size=12, color=SOFT_WHITE, align=PP_ALIGN.CENTER)
    notes(s, "Close warm. Invite questions. Pause and listen.")


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------
def build(images):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    TOTAL = 13

    slide_title(prs, images)
    slide_intro(prs, images, 2, TOTAL)
    slide_agenda(prs, 3, TOTAL)

    # Q: How do you measure success in your current role?
    qa_columns(
        prs, idx=4, total=TOTAL, accent=TEAL,
        eyebrow="Question 1 · Measuring success",
        title="How I measure success in my current role",
        subtitle="I look at three layers — and all three have to hold.",
        columns=[
            ("Business outcome", [
                "Did the actual decision or cycle get better?",
                "Rate-design iteration cycle shortened ~20%.",
                "KPI reporting turnaround cut 45% for 500+ stakeholders.",
                "The metric the business cares about — not a model score in a vacuum.",
            ], TEAL),
            ("System health", [
                "Latency, uptime, and cost in production.",
                "35% end-to-end latency cut on the chatbot.",
                "99.9% uptime on the GCP AI platform.",
                "If it's slow, down, or expensive, it isn't really shipped.",
            ], NAVY),
            ("Adoption & trust", [
                "Are real people using it — and trusting the output?",
                "In daily use across 10 US states.",
                "Teams self-serve; no more engineering tickets.",
                "Every answer is grounded and checked, so people rely on it.",
            ], CORAL),
        ],
        bottom_note="My rule: a model nobody uses or trusts is a failed model — even at high accuracy.",
    )

    # Q: Backtesting recommendation models — best metrics
    qa_columns(
        prs, idx=5, total=TOTAL, accent=NAVY,
        eyebrow="Question 2 · Recommendation models",
        title="Backtesting recommendations — the metrics I rely on",
        subtitle="A recommendation is a ranked list, so I measure ranking — and I backtest honestly.",
        columns=[
            ("Metrics I trust (it's a ranking problem)", [
                "Precision@K / Recall@K — of the top K I show, how many are right.",
                "NDCG / MAP — reward putting the right item near the top, not just in the list.",
                "MRR — how high is the first good hit.",
                "Coverage, diversity, novelty — guard against only pushing popular items.",
                "Calibration — are the scores believable, not just well-ordered.",
            ], NAVY),
            ("How I backtest (so the number is trustworthy)", [
                "Out-of-time split — train on the past, test on the future. Never random shuffle (leakage).",
                "Replay the real decision moment — score with only the data you'd have had then.",
                "Off-policy / counterfactual eval (IPS, doubly-robust) — logged clicks are biased by the OLD recommender.",
                "Segment the results — don't let a strong average hide a weak segment.",
                "Stability over time — a metric that swings week to week isn't reliable.",
            ], TEAL),
        ],
        bottom_note="North star: the backtest is the gate; an online A/B test is the truth. The backtest decides what earns live traffic.",
    )

    # Q: Use of unstructured data to develop features
    qa_columns(
        prs, idx=6, total=TOTAL, accent=CORAL,
        eyebrow="Question 3 · Unstructured data",
        title="Turning unstructured data into model features",
        subtitle="The goal: turn text and documents into reliable, validated, monitored signals.",
        columns=[
            ("Where the signal hides", [
                "PDFs, manuals, contracts, notes.",
                "Support tickets, emails, chat logs.",
                "Spreadsheets that aren't really tabular.",
                "Images and scanned forms.",
            ], CORAL),
            ("How I extract features", [
                "LLM extraction into validated fields (Pydantic schemas).",
                "Embeddings as features for similarity and clustering.",
                "Entity, topic, and sentiment extraction.",
                "RAG-derived signals — 'does a document support X?'",
            ], GOLD),
            ("Real examples from my work", [
                "HOMELOAD: messy utility files (PDF/XML/CSV) → 7 clean columns.",
                "Chatbot Document Agent: pulls answers out of F-150 manuals.",
                "Multi-table parser: turns chaotic Excel into structured tables.",
                "Always validated, versioned, and watched for drift.",
            ], NAVY),
        ],
    )

    # Q: Deciding what features are worth the effort
    qa_columns(
        prs, idx=7, total=TOTAL, accent=GOLD,
        eyebrow="Question 4 · Choosing what to build",
        title="Deciding which features are worth the effort",
        subtitle="I weigh value against the true cost — and I test cheaply before I commit.",
        columns=[
            ("Will it move the metric?", [
                "Run a cheap offline test first — does it lift the backtest metric?",
                "Check feature importance / ablation on a quick prototype.",
                "If it doesn't move the number, it doesn't get built.",
            ], GOLD),
            ("What's the true cost?", [
                "Cost to build AND to maintain — maintenance is the real bill.",
                "Is the data available, fresh, and easy to pipe in?",
                "Any leakage risk or compliance concern?",
            ], CORAL),
            ("Is the source reliable?", [
                "Will it drift or break? Prefer stable, owned sources.",
                "High-impact / low-cost features go first.",
                "Kill features that stop earning their keep.",
            ], TEAL),
        ],
        bottom_note="Bias to a quick, cheap experiment over a long debate — let the data decide what's worth a production pipeline.",
    )

    # Q: Hard conversations with VIP stakeholders
    qa_rows(
        prs, idx=8, total=TOTAL, accent=PLUM,
        eyebrow="Question 5 · VIP stakeholders",
        title="Hard conversations — and keeping the relationship",
        subtitle="Disagree with the idea, never the person. Make it about the shared goal.",
        rows=[
            ("Start from the shared goal", "not my position. We both want the same outcome — frame it that way."),
            ("Bring data and options", "not just a problem. Show trade-offs in their terms: risk, cost, time."),
            ("Be direct, but private", "disagree with the idea openly; protect the person's standing."),
            ("Listen first", "find the real concern behind the ask before I respond to it."),
            ("Say 'no' with a path", "'not that — but here's what I'd recommend instead, and why.'"),
            ("Follow up in writing", "keep them informed so there are never any surprises."),
        ],
    )

    # Q: Once you collect feedback, how do you prioritize?
    qa_columns(
        prs, idx=9, total=TOTAL, accent=TEAL,
        eyebrow="Question 6 · Prioritizing feedback",
        title="Turning collected feedback into a ranked plan",
        subtitle="The loudest voice isn't automatically the top priority.",
        columns=[
            ("Cluster, don't react", [
                "Group feedback into themes first.",
                "One loud request ≠ a priority.",
                "Look for the pattern across many users.",
            ], TEAL),
            ("Score it", [
                "Impact × Reach × Confidence ÷ Effort (RICE-style).",
                "Separate 'must-fix' (broken or trust) from 'nice-to-have.'",
                "Tie every item to the metric that matters.",
            ], NAVY),
            ("Align & close the loop", [
                "Agree the ranked list with stakeholders — no surprises.",
                "Ship in thin slices, re-check after each.",
                "Tell people what I did with their feedback.",
            ], CORAL),
        ],
        bottom_note="Closing the loop is the part people forget — it's what makes them keep giving you honest feedback.",
    )

    # Q: Working style / team configuration / role
    qa_columns(
        prs, idx=10, total=TOTAL, accent=NAVY,
        eyebrow="Question 7 · Working style",
        title="My working style and the role I take",
        subtitle="I default to ownership and over-communicate.",
        columns=[
            ("Team setup I like", [
                "Small, cross-functional pods — data, product, engineering together.",
                "I like owning a problem end-to-end, not a handoff chain.",
                "Tight feedback loops with the people who use the thing.",
            ], NAVY),
            ("The role I take", [
                "The bridge — deep technically, and still fluent with the business.",
                "Often the 'glue' connecting research and production.",
                "Comfortable leading a workstream or contributing heads-down as an IC.",
            ], TEAL),
            ("How I operate", [
                "Write short design docs so decisions are visible.",
                "Demo early and often — show, don't tell.",
                "Over-communicate; surprises are the enemy of trust.",
            ], GOLD),
        ],
    )

    # Q: Agile or waterfall
    qa_columns(
        prs, idx=11, total=TOTAL, accent=CORAL,
        eyebrow="Question 8 · Agile vs waterfall",
        title="How I work — and helping a team move to agile",
        subtitle="I work in agile today, I've done both, and AI work especially rewards short loops.",
        columns=[
            ("How I work today — agile", [
                "Two-week sprints, standups, sprint demos, retros at Ford.",
                "Lightweight backlog grooming and estimates.",
                "Ship in thin slices to get feedback fast.",
            ], CORAL),
            ("Why agile fits AI especially", [
                "AI is uncertain — you can't fully plan it upfront.",
                "Short loops + demos surface data issues and drift early.",
                "Faster to double down on what works and kill what doesn't.",
            ], NAVY),
            ("Helping waterfall → agile", [
                "Start with the rituals that create fast feedback: demos + retros.",
                "Keep planning lightweight; don't cargo-cult ceremonies.",
                "I've worked both ways — happy to help that shift land.",
            ], TEAL),
        ],
        bottom_note="The point of agile isn't the ceremonies — it's shortening the loop between building something and learning if it worked.",
    )

    # Questions for them
    qa_rows(
        prs, idx=12, total=TOTAL, accent=GOLD,
        eyebrow="Questions for you",
        title="A few things I'd love to ask",
        subtitle="To understand how your team works day to day.",
        rows=[
            ("Success in this role", "what does a great first 3-6 months look like to you?"),
            ("The team today", "what does your data/AI stack and workflow look like right now?"),
            ("Recommendations", "how do you measure recommendation quality today — offline, online, or both?"),
            ("Stakeholders", "who are the key stakeholders, and how does the team partner with them?"),
            ("The agile move", "what's driving the shift to agile, and where's the friction so far?"),
        ],
    )

    slide_thanks(prs, images)
    return prs


def main():
    print("Loading images...")
    images = load_images()
    print(f"Images: {len(images)}")
    prs = build(images)
    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
