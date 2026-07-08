#!/usr/bin/env python3
"""Plain-language version of the Charlotte Wang & Kate Landman interview deck.

Same questions and structure as build_charlotte_kate_interview.py, but every
bullet is rewritten as a clear, complete sentence in simple (still professional)
language — so the points are easy to read out loud and easy to follow.

Questions answered:
  1. General introduction
  2. How I measure success in my current role
  3. Backtesting recommendation models — the metrics I rely on
  4. Turning unstructured data into model features
  5. Deciding which features are worth the effort
  6. Hard conversations with VIP stakeholders (and keeping the relationship)
  7. Prioritizing feedback once it's collected
  8. Working style / team configuration / the role I take
  9. Agile vs waterfall (and helping a team make the move)
"""

from __future__ import annotations

from pathlib import Path

import requests
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "presentation-assets"
OUTPUT = ROOT / "Moein-Razavi-Interview-Charlotte-Kate-Simple.pptx"

# ---------------------------------------------------------------------------
# Palette (matches the other decks)
# ---------------------------------------------------------------------------
NAVY = RGBColor(20, 42, 78)
NAVY_DARK = RGBColor(13, 27, 51)
INK = RGBColor(28, 36, 54)
BODY = RGBColor(70, 82, 102)
MUTED = RGBColor(130, 140, 158)
LINE = RGBColor(220, 226, 236)
CREAM = RGBColor(250, 248, 244)
CORAL = RGBColor(232, 96, 80)
GOLD = RGBColor(232, 178, 70)
TEAL = RGBColor(52, 152, 162)
PLUM = RGBColor(150, 90, 158)
WHITE = RGBColor(255, 255, 255)
SOFT_WHITE = RGBColor(238, 244, 252)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

IMAGE_URLS = {
    "hero": "https://images.unsplash.com/photo-1454165804606-c3d57bc86b40?w=1920&q=80",
    "team": "https://images.unsplash.com/photo-1522071820081-009f0129c71c?w=1920&q=80",
}


def load_images() -> dict[str, Path]:
    ASSETS.mkdir(exist_ok=True)
    paths: dict[str, Path] = {}
    s = requests.Session()
    s.headers["User-Agent"] = "Moein-Interview-Deck/1.0"
    for key, url in IMAGE_URLS.items():
        dest = ASSETS / f"{key}.jpg"
        if not dest.exists():
            try:
                r = s.get(url, timeout=30)
                r.raise_for_status()
                dest.write_bytes(r.content)
            except Exception as exc:
                print(f"[warn] {key}: {exc}")
                continue
        paths[key] = dest
    return paths


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def fill(shape, color, transparency=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if transparency is not None:
        shape.fill.transparency = transparency


def rect(slide, l, t, w, h, color, transparency=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color, transparency)
    return s


def round_rect(slide, l, t, w, h, color, corner=0.05, transparency=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color, transparency)
    try:
        s.adjustments[0] = corner
    except Exception:
        pass
    return s


def text(slide, l, t, w, h, content, *, size=14, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=None, italic=False, font=None, wrap=True):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = content
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    if font:
        r.font.name = font
    return box


def bullets(slide, l, t, w, h, items, *, size=12, color=BODY,
            bullet=CORAL, space=7):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        r1 = p.add_run()
        r1.text = "•  "
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = bullet
        r2 = p.add_run()
        r2.text = item
        r2.font.size = Pt(size)
        r2.font.color.rgb = color


def image_overlay(slide, img: Path, transparency=0.55, color=NAVY_DARK):
    slide.shapes.add_picture(str(img), 0, 0, width=SLIDE_W, height=SLIDE_H)
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, color, transparency=transparency)


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


# ---------------------------------------------------------------------------
# Page chrome
# ---------------------------------------------------------------------------
def page(prs, *, title, subtitle=None, eyebrow=None, eyebrow_color=TEAL):
    s = blank(prs)
    bg(s, CREAM)
    rect(s, Inches(0.7), Inches(0.6), Inches(0.6), Inches(0.07), eyebrow_color)
    if eyebrow:
        text(s, Inches(0.7), Inches(0.72), Inches(11.9), Inches(0.3),
             eyebrow.upper(), size=11, bold=True, color=eyebrow_color)
        t_top = Inches(1.02)
    else:
        t_top = Inches(0.78)
    text(s, Inches(0.7), t_top, Inches(12), Inches(0.9),
         title, size=30, bold=True, color=NAVY)
    if subtitle:
        text(s, Inches(0.7), t_top + Inches(0.88), Inches(12), Inches(0.45),
             subtitle, size=14, color=BODY, italic=True)
    return s


def footer(slide, idx, total):
    rect(slide, Inches(0.7), Inches(7.05), Inches(11.93), Inches(0.012), LINE)
    text(slide, Inches(0.7), Inches(7.13), Inches(9), Inches(0.25),
         "Moein Razavi  ·  Interview with Charlotte Wang & Kate Landman",
         size=9, color=MUTED)
    text(slide, Inches(11.4), Inches(7.13), Inches(1.4), Inches(0.25),
         f"{idx} / {total}", size=9, color=MUTED, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# Reusable answer layouts
# ---------------------------------------------------------------------------
def value_card(s, x, y, w, h, header, items, accent,
               *, header_size=14, body_size=11.5):
    round_rect(s, x, y, w, h, WHITE, corner=0.05)
    rect(s, x, y, w, Inches(0.08), accent)
    text(s, x + Inches(0.25), y + Inches(0.24), w - Inches(0.5), Inches(0.55),
         header, size=header_size, bold=True, color=NAVY)
    bullets(s, x + Inches(0.25), y + Inches(0.92), w - Inches(0.5),
            h - Inches(1.05), items, size=body_size, color=BODY,
            bullet=accent, space=7)


def qa_columns(prs, *, idx, total, eyebrow, title, subtitle, columns, accent,
               bottom_note=None, body_size=11.5):
    """columns: list of (header, [bullets], accent_color)."""
    s = page(prs, eyebrow=eyebrow, eyebrow_color=accent, title=title,
             subtitle=subtitle)
    n = len(columns)
    gap = Inches(0.25)
    col_w = (Inches(11.93) - gap * (n - 1)) / n
    col_top = Inches(2.45)
    col_h = Inches(3.7) if bottom_note else Inches(4.35)
    for i, (header, items, caccent) in enumerate(columns):
        x = Inches(0.7) + i * (col_w + gap)
        value_card(s, x, col_top, col_w, col_h, header, items, caccent,
                   body_size=body_size)
    if bottom_note:
        ny = col_top + col_h + Inches(0.18)
        round_rect(s, Inches(0.7), ny, Inches(11.93), Inches(0.62), NAVY,
                   corner=0.08)
        text(s, Inches(0.95), ny, Inches(11.4), Inches(0.62),
             bottom_note, size=12, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    footer(s, idx, total)
    return s


def qa_rows(prs, *, idx, total, eyebrow, title, subtitle, rows, accent):
    """rows: list of (lead, detail)."""
    s = page(prs, eyebrow=eyebrow, eyebrow_color=accent, title=title,
             subtitle=subtitle)
    top = Inches(2.5)
    row_h = Inches(0.6)
    gap = Inches(0.1)
    for i, (lead, detail) in enumerate(rows):
        y = top + i * (row_h + gap)
        round_rect(s, Inches(0.7), y, Inches(11.93), row_h, WHITE, corner=0.06)
        rect(s, Inches(0.7), y, Inches(0.6), row_h, accent)
        text(s, Inches(0.7), y, Inches(0.6), row_h, str(i + 1),
             size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        box = s.shapes.add_textbox(Inches(1.55), y, Inches(10.9), row_h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{lead}  "
        r1.font.size = Pt(12.5); r1.font.bold = True; r1.font.color.rgb = NAVY
        r2 = p.add_run()
        r2.text = f"— {detail}"
        r2.font.size = Pt(12); r2.font.color.rgb = BODY
    footer(s, idx, total)
    return s


# ---------------------------------------------------------------------------
# Bespoke slides
# ---------------------------------------------------------------------------
def slide_title(prs, images):
    s = blank(prs)
    if "hero" in images:
        image_overlay(s, images["hero"], transparency=0.58, color=NAVY_DARK)
    else:
        bg(s, NAVY_DARK)

    rect(s, 0, Inches(2.55), Inches(0.45), Inches(0.18), TEAL)
    text(s, Inches(0.85), Inches(2.4), Inches(11.5), Inches(0.4),
         "INTERVIEW  ·  PREPARED FOR CHARLOTTE WANG & KATE LANDMAN",
         size=13, bold=True, color=GOLD)
    text(s, Inches(0.85), Inches(2.95), Inches(11.5), Inches(1.3),
         "Hi, I'm Moein.", size=56, bold=True, color=WHITE)
    text(s, Inches(0.85), Inches(4.25), Inches(11.7), Inches(0.6),
         "I'm a senior AI engineer who builds tools people actually use — and works closely with them.",
         size=19, color=SOFT_WHITE)
    text(s, Inches(0.85), Inches(5.05), Inches(11.7), Inches(0.45),
         "Here's how I think about success, data, working with people, and how I like to work.",
         size=15, italic=True, color=GOLD)
    text(s, Inches(0.85), Inches(6.5), Inches(11.5), Inches(0.35),
         "razavi.moein94@gmail.com  ·  linkedin.com/in/moein-razavi",
         size=11, color=SOFT_WHITE)
    notes(s, "Warm open. Thank Charlotte and Kate. This round is about how I work, "
             "so I'll keep examples concrete and let them steer.")


def slide_intro(prs, images, idx, total):
    s = page(prs, eyebrow="General introduction",
             title="A little about me", eyebrow_color=TEAL,
             subtitle="The short version, in plain language.")
    points = [
        "I'm a senior AI/ML engineer. I have a Ph.D. from Texas A&M, where I finished with a perfect 4.0 GPA.",
        "For the last 5+ years I've built AI that real companies use every day. Right now I'm at Ford.",
        "I don't just build a model and hand it off — I own the whole thing: the business problem, the model, and keeping it running well in production.",
        "I'm comfortable in both worlds: I can do the deep technical work, and I can explain it clearly to people who aren't technical.",
        "Three things I care about: making a real, measurable difference; earning people's trust; and actually shipping, not just demoing.",
    ]
    bullets(s, Inches(0.7), Inches(2.45), Inches(7.5), Inches(4.3),
            points, size=15, color=INK, bullet=TEAL, space=13)
    if "team" in images:
        s.shapes.add_picture(str(images["team"]),
                             Inches(8.6), Inches(2.6),
                             width=Inches(4.1), height=Inches(3.9))
    footer(s, idx, total)
    notes(s, "Keep it to ~60 seconds. Land the 'bridge between technical and business' point.")


def slide_agenda(prs, idx, total):
    s = page(prs, eyebrow="What I'd love to cover",
             title="The questions you shared — my take on each",
             eyebrow_color=TEAL,
             subtitle="Please stop me on any of them — I'm happy to go deeper with real examples.")
    topics = [
        ("1", "Measuring success", "How I know my work is actually helping", TEAL),
        ("2", "Testing recommendation models", "How I test recommendations fairly on past data", NAVY),
        ("3", "Using messy data", "Turning text and documents into useful inputs", CORAL),
        ("4", "Choosing what to build", "Deciding which new features are worth it", GOLD),
        ("5", "Hard conversations", "Disagreeing with senior people without losing trust", PLUM),
        ("6", "Prioritizing feedback", "Turning lots of feedback into a clear plan", TEAL),
        ("7", "Working style", "How I work with a team", NAVY),
        ("8", "Agile vs waterfall", "How I work — and helping the move to agile", CORAL),
    ]
    gap_x = Inches(0.25)
    gap_y = Inches(0.16)
    card_w = (Inches(11.93) - gap_x) / 2
    card_h = Inches(0.95)
    top = Inches(2.45)
    for i, (num, title, desc, accent) in enumerate(topics):
        col = i % 2
        row = i // 2
        x = Inches(0.7) + col * (card_w + gap_x)
        y = top + row * (card_h + gap_y)
        round_rect(s, x, y, card_w, card_h, WHITE, corner=0.07)
        rect(s, x, y, Inches(0.6), card_h, accent)
        text(s, x, y, Inches(0.6), card_h, num, size=20, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + Inches(0.8), y + Inches(0.16), card_w - Inches(1.0), Inches(0.4),
             title, size=14, bold=True, color=NAVY)
        text(s, x + Inches(0.8), y + Inches(0.52), card_w - Inches(1.0), Inches(0.38),
             desc, size=11, color=BODY)
    footer(s, idx, total)


def slide_thanks(prs, images):
    s = blank(prs)
    if "hero" in images:
        image_overlay(s, images["hero"], transparency=0.62, color=NAVY_DARK)
    else:
        bg(s, NAVY_DARK)
    rect(s, Inches(6.0), Inches(2.0), Inches(1.3), Inches(0.18), TEAL)
    text(s, 0, Inches(2.4), SLIDE_W, Inches(1.2),
         "Thank you.", size=58, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, 0, Inches(3.7), SLIDE_W, Inches(0.5),
         "I'd love to hear how your team works — and answer anything.",
         size=18, color=GOLD, align=PP_ALIGN.CENTER)
    card_w = Inches(8.0)
    cx = (SLIDE_W - card_w) / 2
    round_rect(s, cx, Inches(4.7), card_w, Inches(1.85), NAVY_DARK, corner=0.04)
    text(s, cx, Inches(4.88), card_w, Inches(0.4), "MOEIN RAZAVI",
         size=12, bold=True, color=GOLD, align=PP_ALIGN.CENTER)
    text(s, cx, Inches(5.25), card_w, Inches(0.4),
         "razavi.moein94@gmail.com", size=14, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, cx, Inches(5.65), card_w, Inches(0.4),
         "(979) 676-7486", size=13, color=SOFT_WHITE, align=PP_ALIGN.CENTER)
    text(s, cx, Inches(6.05), card_w, Inches(0.4),
         "linkedin.com/in/moein-razavi  ·  github.com/moeinrazavi",
         size=12, color=SOFT_WHITE, align=PP_ALIGN.CENTER)
    notes(s, "Close warm. Invite questions. Pause and listen.")


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------
def build(images):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    TOTAL = 13

    slide_title(prs, images)
    slide_intro(prs, images, 2, TOTAL)
    slide_agenda(prs, 3, TOTAL)

    # Q1: How do you measure success in your current role?
    qa_columns(
        prs, idx=4, total=TOTAL, accent=TEAL,
        eyebrow="Question 1 · Measuring success",
        title="How I measure success in my current role",
        subtitle="I look at three things, and all three have to be true.",
        columns=[
            ("Did the business get better?", [
                "My first question is simple: did the business actually improve because of my work?",
                "Example: I helped cut the time to design a new pricing plan by about 20%.",
                "Example: I made KPI reports come out 45% faster for over 500 people who rely on them.",
            ], TEAL),
            ("Is the system healthy?", [
                "I watch three things: speed, reliability, and cost.",
                "I cut the chatbot's response time by 35%.",
                "My AI platform on Google Cloud stays up 99.9% of the time.",
                "If a tool is slow, keeps breaking, or costs too much, it isn't really finished.",
            ], NAVY),
            ("Do people use and trust it?", [
                "Are real people using it — and do they trust what it tells them?",
                "My chatbot is used every day by teams in 10 US states.",
                "People answer their own questions now, instead of waiting on a ticket.",
                "Every answer is based on real data and double-checked, so people feel safe relying on it.",
            ], CORAL),
        ],
        bottom_note="My simple rule: if no one uses a model or trusts it, it has failed — even if it is very accurate.",
    )

    # Q2: Backtesting recommendation models — best metrics
    qa_columns(
        prs, idx=5, total=TOTAL, accent=NAVY,
        eyebrow="Question 2 · Recommendation models",
        title="How I test recommendation models on past data",
        subtitle="A recommendation is just a ranked list of suggestions, so I measure how good that order is — and I test it honestly.",
        columns=[
            ("Which numbers I look at", [
                "Did the top suggestions turn out to be the right ones? (Precision and Recall)",
                "Are the best items near the top of the list, not buried lower down? (NDCG and MAP)",
                "How far down the list is the first good suggestion? (MRR)",
                "Is the list varied, or am I just showing everyone the same popular items? (coverage and diversity)",
            ], NAVY),
            ("How I test it fairly", [
                "I train on older data and test on newer data — never mix them, or the test looks better than real life.",
                "I only let the model use the information it would have actually had at that moment in time.",
                "Past clicks are biased because they came from the old system, so I correct for that (off-policy testing).",
                "I check each customer group on its own, so a good overall score can't hide a group it's failing.",
            ], TEAL),
        ],
        bottom_note="Bottom line: testing on past data tells me what's promising, but a live A/B test is the real proof. The past-data test decides what's worth trying live.",
    )

    # Q3: Use of unstructured data to develop features
    qa_columns(
        prs, idx=6, total=TOTAL, accent=CORAL,
        eyebrow="Question 3 · Using messy data",
        title="Turning text and documents into useful model inputs",
        subtitle="'Unstructured data' just means messy information like text and documents. The goal is to turn it into clean, reliable inputs a model can use.",
        columns=[
            ("Where the useful information hides", [
                "PDF files, manuals, and contracts.",
                "Support tickets, emails, and chat conversations.",
                "Spreadsheets that look like tables but really aren't.",
                "Photos and scanned forms.",
            ], CORAL),
            ("How I pull it out", [
                "I use an AI model to read the text and pull out specific facts into clean, fixed fields.",
                "I turn text into numbers so the model can compare and group similar things (embeddings).",
                "I pull out names, topics, and whether the tone is positive or negative.",
                "I ask 'does this document support a certain claim?' and use the answer as a signal.",
            ], GOLD),
            ("Real examples from my work", [
                "HOMELOAD: I turned messy utility files (PDF, XML, CSV) into 7 clean, standard columns.",
                "Chatbot: one part reads Ford F-150 manuals and pulls out the exact answer.",
                "A tool I built turns chaotic Excel sheets into clean, usable tables.",
                "In every case I check the data, save versions, and watch it for changes over time.",
            ], NAVY),
        ],
    )

    # Q4: Deciding what features are worth the effort
    qa_columns(
        prs, idx=7, total=TOTAL, accent=GOLD,
        eyebrow="Question 4 · Choosing what to build",
        title="Deciding which new features are worth the effort",
        subtitle="Before building a feature, I weigh how much it helps against how much it truly costs — and I run a small test first.",
        columns=[
            ("Will it actually help?", [
                "First I run a quick, cheap test to see if the feature really improves the results.",
                "I check how much the model leans on it before building anything big.",
                "If it doesn't improve the result, I don't build it.",
            ], GOLD),
            ("What does it really cost?", [
                "I count the cost to build it AND to keep it running — upkeep is the bill people forget.",
                "Is the data easy to get, fresh, and dependable?",
                "Could it leak information or create a compliance problem?",
            ], CORAL),
            ("Can I rely on the source?", [
                "Will the data source change or break later? I prefer sources we control.",
                "I build the high-impact, low-effort features first.",
                "If a feature stops being useful, I remove it.",
            ], TEAL),
        ],
        bottom_note="I'd rather run a quick, cheap experiment than have a long debate — let the data decide what deserves a full pipeline.",
    )

    # Q5: Hard conversations with VIP stakeholders
    qa_rows(
        prs, idx=8, total=TOTAL, accent=PLUM,
        eyebrow="Question 5 · Senior stakeholders",
        title="Hard conversations — without damaging the relationship",
        subtitle="I disagree with the idea, never the person, and I keep it about the goal we both share.",
        rows=[
            ("Start with the shared goal", "we both want the same outcome, so I begin there instead of defending my side."),
            ("Bring data and choices", "I don't just bring a problem — I bring options and explain the trade-offs in their terms: risk, cost, and time."),
            ("Be honest, but private", "I'll disagree with an idea openly, but I never make the person look bad."),
            ("Listen first", "I find out the real worry behind their request before I respond to it."),
            ("Say 'no' with a plan", "instead of just 'no,' I say 'not that — here's what I'd suggest instead, and why.'"),
            ("Put it in writing", "I follow up in writing so they stay informed and are never surprised."),
        ],
    )

    # Q6: Once you collect feedback, how do you prioritize?
    qa_columns(
        prs, idx=9, total=TOTAL, accent=TEAL,
        eyebrow="Question 6 · Prioritizing feedback",
        title="Turning a pile of feedback into a clear, ranked plan",
        subtitle="The loudest person isn't automatically right, so I work through it in a simple way.",
        columns=[
            ("Group it, don't react", [
                "First I group similar feedback into themes.",
                "One loud request doesn't make it a priority.",
                "I look for the pattern that many people are pointing at.",
            ], TEAL),
            ("Score each item", [
                "I score each item by how much it helps, how many people it affects, how sure I am, and how much work it takes.",
                "I separate 'must fix' problems from 'nice to have' wishes.",
                "I tie every item back to the goal that matters most.",
            ], NAVY),
            ("Agree and close the loop", [
                "I agree on the ranked list with stakeholders, so there are no surprises.",
                "I ship small pieces and check again after each one.",
                "I tell people what I did with their feedback — the step most people skip.",
            ], CORAL),
        ],
        bottom_note="Telling people what happened with their feedback is what makes them keep giving you honest input.",
    )

    # Q7: Working style / team configuration / role
    qa_columns(
        prs, idx=10, total=TOTAL, accent=NAVY,
        eyebrow="Question 7 · Working style",
        title="How I work, and the role I usually take",
        subtitle="I take ownership by default, and I keep everyone in the loop.",
        columns=[
            ("The team setup I like", [
                "Small teams where data, product, and engineering work side by side.",
                "I like owning a problem from start to finish, not passing it down a chain.",
                "Quick, regular feedback from the people who actually use what we build.",
            ], NAVY),
            ("The role I take", [
                "I'm the bridge: strong on the technical side, and comfortable talking to the business.",
                "I'm often the 'glue' that connects the research and the real product.",
                "I'm happy to lead a piece of work, or to put my head down and build as part of the team.",
            ], TEAL),
            ("How I operate day to day", [
                "I write short documents so everyone can see why we made each decision.",
                "I show working demos early and often, instead of just talking about it.",
                "I over-communicate, because surprises break trust.",
            ], GOLD),
        ],
    )

    # Q8: Agile or waterfall
    qa_columns(
        prs, idx=11, total=TOTAL, accent=CORAL,
        eyebrow="Question 8 · Agile vs waterfall",
        title="How I work, and helping a team move to agile",
        subtitle="I work in agile today, I've used both, and AI work especially benefits from short, fast cycles.",
        columns=[
            ("How I work today — agile", [
                "At Ford I work in two-week sprints, with daily standups, demos, and retros.",
                "We keep planning light and estimate the work as a team.",
                "We release in small pieces so we get feedback quickly.",
            ], CORAL),
            ("Why agile fits AI work", [
                "AI is uncertain — you can't fully plan it on paper up front.",
                "Short cycles and demos catch data problems early.",
                "It's easier to keep what works and drop what doesn't.",
            ], NAVY),
            ("Helping a team make the move", [
                "I'd start with the habits that give the fastest feedback: regular demos and retros.",
                "Keep the planning simple — don't copy rituals just for the sake of it.",
                "I've worked both ways, so I can help make the change stick.",
            ], TEAL),
        ],
        bottom_note="Agile isn't really about the meetings — it's about shortening the time between building something and learning whether it worked.",
    )

    # Questions for them
    qa_rows(
        prs, idx=12, total=TOTAL, accent=GOLD,
        eyebrow="Questions for you",
        title="A few things I'd love to ask",
        subtitle="To understand how your team works day to day.",
        rows=[
            ("Success in this role", "what would a great first 3 to 6 months look like to you?"),
            ("Your team today", "what do your data and AI tools and your daily workflow look like right now?"),
            ("Recommendations", "how do you measure whether a recommendation is good today — on past data, live, or both?"),
            ("Stakeholders", "who are the main people you work with, and how does the team partner with them?"),
            ("The move to agile", "what's pushing the change to agile, and where has it been bumpy so far?"),
        ],
    )

    slide_thanks(prs, images)
    return prs


def main():
    print("Loading images...")
    images = load_images()
    print(f"Images: {len(images)}")
    prs = build(images)
    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
