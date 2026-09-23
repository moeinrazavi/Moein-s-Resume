"""Build role-specific resumes from existing records and user corrections.

Employment titles, CereVu contract status, and graduate degrees were confirmed
by the user. Dates otherwise follow the latest existing resume. Technical
implementation details use build_simple_portfolio.py and the existing .tex
resumes; these are source descriptions, not independently audited code.

Public research cross-checks:
  https://arxiv.org/abs/2309.11097 (54 participants, 40 days)
  https://github.com/TAMUCogLab/OpenSync (LSL, not WebSocket)
  https://arxiv.org/abs/2107.14367
  https://arxiv.org/pdf/2107.14367v2 (Table 5: RMS time lag 50.69-65.28 us
  across four device configurations in 5- and 60-minute experiments; includes
  device and recording-system effects, not a per-sample worst-case guarantee)

Quantification retained from source descriptions: five file formats and seven
output columns (portfolio HOMELOAD story); two iOS apps and nine physiological
signals (multiple existing resumes); approximately 3,000 simulated scenarios
per dispatch cycle (portfolio V2G story). These describe implementation scope,
not independently verified business outcomes or deployment scale.

Excluded: unsubstantiated business/performance percentages, production scale,
clinical deployment claims, publication/citation totals, honorary titles,
unverified certifications, and broad technology lists without project context.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output' / 'resume'


class Resume:
    def __init__(self, slug, headline, summary):
        self.slug = slug
        self.doc = Document()
        self.text = []
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.7)
        normal = self.doc.styles['Normal']
        normal.font.name = 'Arial'
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.04
        for name, size in [('Title', 22), ('Heading 1', 11.5), ('Heading 2', 10.5)]:
            style = self.doc.styles[name]
            style.font.name = 'Arial'
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_before = Pt(10 if name == 'Heading 1' else 5)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.keep_with_next = True
        self.doc.styles['Title'].paragraph_format.space_before = Pt(0)
        for style in self.doc.styles:
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        for level in self.doc.part.numbering_part.element.iter(qn('w:lvl')):
            fmt = level.find(qn('w:numFmt'))
            if fmt is not None and fmt.get(qn('w:val')) == 'bullet':
                level.find(qn('w:lvlText')).set(qn('w:val'), '\u2022')
                props = level.find(qn('w:rPr'))
                if props is not None:
                    fonts = props.find(qn('w:rFonts'))
                    if fonts is not None:
                        fonts.set(qn('w:ascii'), 'Arial')
                        fonts.set(qn('w:hAnsi'), 'Arial')
        self.doc.core_properties.author = 'Moein Razavi'
        self.doc.core_properties.title = f'Moein Razavi - {headline}'
        self.doc.core_properties.subject = headline
        self.doc.core_properties.keywords = ''
        self.p('Moein Razavi', 'Title')
        self.p(headline, bold=True)
        self.p('McKinney, TX | (979) 676-7486 | razavi.moein94@gmail.com')
        p = self.doc.add_paragraph()
        self.link(p, 'linkedin.com/in/moein-razavi', 'https://linkedin.com/in/moein-razavi')
        p.add_run(' | ')
        self.link(p, 'github.com/moeinrazavi', 'https://github.com/moeinrazavi')
        self.text.append('linkedin.com/in/moein-razavi | github.com/moeinrazavi')
        self.heading('Professional Summary')
        self.p(summary)

    def link(self, p, label, url):
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '000000')
        props.append(color)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = label
        run.append(text)
        hyperlink.append(run)
        p._p.append(hyperlink)

    def p(self, text, style=None, bold=None, prefix=None):
        p = self.doc.add_paragraph(style=style)
        if prefix:
            p.add_run(prefix).bold = True
            p.add_run(text[len(prefix):])
        else:
            run = p.add_run(text)
            if bold is not None:
                run.bold = bold
        p.paragraph_format.keep_together = True
        self.text.append(text)
        return p

    def heading(self, text, new_page=False):
        self.text.append('')
        p = self.p(text, 'Heading 1')
        p.paragraph_format.page_break_before = new_page

    def role(self, employer, title, dates):
        self.p(employer, 'Heading 2')
        p = self.p(f'{title} | {dates}')
        p.paragraph_format.keep_with_next = True

    def bullet(self, text, prefix=None):
        p = self.p(text, 'List Bullet', prefix=prefix)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        self.text[-1] = '- ' + self.text[-1]

    def skills(self, rows):
        self.heading('Technical Skills')
        for label, content in rows:
            self.p(f'{label}: {content}', prefix=label + ':')

    def education(self):
        self.heading('Education')
        self.p('Ph.D., Industrial Engineering | Texas A&M University | Dec 2023', prefix='Ph.D., Industrial Engineering')
        self.p('Master of Engineering, Computer Engineering | Texas A&M University | May 2021', prefix='Master of Engineering, Computer Engineering')

    def save(self):
        OUT.mkdir(parents=True, exist_ok=True)
        assert not self.doc.tables
        self.doc.save(OUT / f'{self.slug}.docx')
        (OUT / f'{self.slug}.txt').write_text('\n'.join(self.text) + '\n')
        print(f'{self.slug}: {len(" ".join(self.text).split())} words')


def ai_resume():
    r = Resume(
        'Moein-Razavi-AI-ML-Data-Scientist-Resume',
        'AI/ML Engineer | Data Scientist',
        'Ph.D.-trained data scientist building machine learning and AI applications for electric vehicle energy analytics. At Ford, develops forecasting models, LLM-based data tools, and cloud services using Python, SQL, and Google Cloud. Combines model development with data engineering and application delivery; published research covers wearable stress detection and sensor synchronization.'
    )
    r.skills([
        ('Languages and Data', 'Python, SQL, pandas, NumPy, PySpark, BigQuery, Apache Airflow'),
        ('Machine Learning', 'PyTorch, TensorFlow, scikit-learn, XGBoost, Hugging Face, Optuna, time-series forecasting'),
        ('Generative AI', 'LangChain, Gemini, retrieval-augmented generation (RAG), text-to-SQL, BM25, LLM routing'),
        ('Cloud and Deployment', 'GCP, Vertex AI, Cloud Run, FastAPI, Flask, Docker, Terraform, Tekton CI/CD'),
    ])
    r.heading('Professional Experience')
    r.role('Ford Motor Company', 'Data Scientist', 'Aug 2023 - Present')
    r.bullet('Built an EV analytics assistant that answers questions from BigQuery data and technical documents. Routed requests to SQL generation or BM25 document search, with a separate LLM check against domain rules.')
    r.bullet('Standardized utility data from 5 file formats (CSV, Excel, PDF, XML, HTML) into a 7-column schema using FastAPI, pandas, and Vertex AI Gemini, with checks on timestamps, value ranges, and required fields.')
    r.bullet('Developed vehicle-to-grid (V2G) dispatch software using vehicle-availability forecasts and approximately 3,000 Monte Carlo scenarios per decision to estimate fleet capacity and select vehicles for energy commitments.')
    r.bullet('Trained transformer forecasting models, including Temporal Fusion Transformer (TFT) and Autoformer, with Optuna tuning on Vertex AI GPUs; compared forecasts with a previous-week baseline.')
    r.bullet('Built XGBoost models to estimate EV electricity bills from vehicle telemetry, utility tariffs, parking behavior, and household consumption, supplying bill predictions to a rate recommendation application.')
    r.bullet('Automated energy-data preparation with PySpark, BigQuery, and Airflow; delivered a Flask/Dash rate recommendation application on Cloud Run with Tekton deployment pipelines.')
    r.role('CereVu Medical, Inc.', 'Software Developer / Research Engineer (Contract)', 'Jun 2023 - Sep 2024')
    r.bullet('Developed 2 SwiftUI apps, Vitality and Clinical, to display 9 physiological signals from wearable sensors, with CoreBluetooth integration and Python services for biosignal streaming.')
    r.bullet('Developed a React/Next.js dashboard with MongoDB and role-based access for reviewing physiological data.')

    r.heading('Research Experience', new_page=True)
    r.role('Texas A&M University', 'Research Assistant', 'Sep 2018 - May 2024')
    r.bullet('Developed and evaluated stress-detection models using heart-rate and hand-acceleration data from 54 college students over 40 days, with time-stamped self-reports as stress labels.')
    r.bullet('Compared machine learning models and identified heart-rate and motion features associated with stress; reported XGBoost findings in a first-author study to inform wearable stress-detection tools.')
    r.bullet('Co-developed OpenSync to synchronize neuroscience sensors using Python and Lab Streaming Layer; published tests measured below 0.07 ms root-mean-square (RMS) recording-time lag across 4 device configurations.')
    r.bullet('Developed computer vision methods for physical-distance and face-mask monitoring; co-authored a scoping review of ML, deep learning, and preprocessing for stress-related mental health research.')

    r.heading('Selected AI Project')
    r.p('Business Action Prioritizer', 'Heading 2')
    r.bullet('Built a Google ADK application that converts uploaded spreadsheets and a scoring rubric into a proposed analysis plan, requests user approval, and executes the approved steps with a PandasAI helper.')
    r.bullet('Generated ranked recommendations and reports, with downloadable outputs stored in Google Cloud Storage through a Model Context Protocol (MCP) tool.')

    r.education()
    r.heading('Selected Publications')
    r.p('Razavi, M., et al. (2024). Machine Learning, Deep Learning, and Data Preprocessing Techniques for Detecting, Predicting, and Monitoring Stress and Stress-Related Mental Disorders: Scoping Review. JMIR Mental Health, 11, e53714. doi:10.2196/53714.')
    r.p('Razavi, M., et al. (2023). Evaluating Mental Stress Among College Students Using Heart Rate and Hand Acceleration Data Collected from Wearable Sensors. arXiv preprint, arXiv:2309.11097.')
    r.p('Razavi, M., et al. (2022). OpenSync: An open-source platform for synchronizing multiple measures in neuroscience experiments. Journal of Neuroscience Methods, 369, 109458. doi:10.1016/j.jneumeth.2021.109458.')
    r.save()


def software_resume():
    r = Resume(
        'Moein-Razavi-Software-Engineer-Resume',
        'Software Engineer | Backend and Mobile Applications',
        'Software developer and data scientist building Python services, cloud data applications, and SwiftUI mobile apps. Experience spans Ford energy analytics, CereVu wearable sensor software, and independently developed inspection and social applications. Works across APIs, data storage, authentication, device integration, and automated releases. Ph.D. in Industrial Engineering and master\'s in Computer Engineering.'
    )
    r.skills([
        ('Languages', 'Python, Swift, TypeScript, JavaScript, SQL'),
        ('Backend and Web', 'FastAPI, Flask, REST APIs, WebSocket, React, Next.js, Node.js'),
        ('Mobile', 'SwiftUI, UIKit, CoreBluetooth, Combine, AVFoundation, MVVM, TestFlight'),
        ('Data and Cloud', 'GCP, Cloud Run, BigQuery, MongoDB, Firebase, Supabase, PySpark, Airflow'),
        ('Delivery and Access', 'Docker, Terraform, Tekton, GitHub Actions, Fastlane, OAuth 2.0, JWT, SAML'),
    ])
    r.heading('Professional Experience')
    r.role('Ford Motor Company', 'Data Scientist', 'Aug 2023 - Present')
    r.bullet('Delivered a utility rate recommendation application that connects energy-data analysis with an authenticated Dash interface, using Flask REST APIs, BigQuery, SAML/ADFS, Cloud Run, and Tekton CI/CD.')
    r.bullet('Built a FastAPI service that converts 5 utility-file formats into a standard 7-column dataset, with validation, retry limits, and BigQuery MERGE operations to support repeatable data loads.')
    r.bullet('Built PySpark and Airflow workflows for energy-simulation data, plus REST APIs and Looker Studio dashboards for operational and EV discharge reporting.')
    r.bullet('Built a code-analysis web application using Python AST parsing and LLM-generated recommendations; streamed each result to the browser through FastAPI so users could review findings while analysis continued.')
    r.role('CereVu Medical, Inc.', 'Software Developer / Research Engineer (Contract)', 'Jun 2023 - Sep 2024')
    r.bullet('Developed 2 iOS apps, Vitality and Clinical, to acquire and chart 9 physiological signals from wearable sensors using SwiftUI, CoreBluetooth, and Swift Charts.')
    r.bullet('Migrated the Clinical app from UIKit to modular SwiftUI components using MVVM-Coordinator, supporting 3 user roles (patient, physician, administrator), sensor calibration, and CSV export.')
    r.bullet('Built Flask/FastAPI services for biosignal streaming over WebSocket and a React/Next.js dashboard backed by MongoDB for reviewing patient data.')
    r.bullet('Implemented authentication and role-based access with OAuth 2.0 and JWT, and supported iOS release workflows through TestFlight.')

    r.heading('Selected Software Projects', new_page=True)
    r.p('Inspecalytics | Field Inspection Application', 'Heading 2')
    r.bullet('Built a SwiftUI inspection app and companion web interface sharing a Supabase backend, connecting field data capture, annotated photos, and report generation.')
    r.bullet('Enabled inspections without a network connection through local data capture and later cloud synchronization, using change detection and concurrent image transfers to reconcile records and photos.')
    r.bullet('Automated Word and PowerPoint reports from inspection data and annotated photos; configured Fastlane and GitHub Actions to build and distribute TestFlight releases.')
    r.p('Clink Social | Social Application', 'Heading 2')
    r.bullet('Built a SwiftUI application using Firebase Authentication, Firestore, Storage, and Cloud Messaging, with real-time messaging, MapKit discovery, and a social feed.')
    r.bullet('Implemented TypeScript/Node.js cloud functions for feed aggregation and push notifications, plus a QR-based cross-device sign-in flow using time-limited codes and token exchange.')

    r.heading('Research Experience')
    r.role('Texas A&M University', 'Research Assistant', 'Sep 2018 - May 2024')
    r.bullet('Co-developed OpenSync, an open-source Python platform that integrates EEG, eye-tracking, physiological sensors, and experiment events using Lab Streaming Layer for synchronized recording.')
    r.bullet('Validated recording timing across 4 device configurations in 5- and 60-minute tests; published benchmarks measured below 0.07 ms root-mean-square (RMS) recording-time lag.')
    r.bullet('Developed GlucoseCoach, an iOS application with a Flask backend for personalized glucose monitoring.')
    p = r.doc.add_paragraph()
    r.link(p, 'OpenSync source: github.com/TAMUCogLab/OpenSync', 'https://github.com/TAMUCogLab/OpenSync')
    r.text.append('OpenSync source: github.com/TAMUCogLab/OpenSync')

    r.education()
    r.heading('Selected Publication')
    r.p('Razavi, M., et al. (2022). OpenSync: An open-source platform for synchronizing multiple measures in neuroscience experiments. Journal of Neuroscience Methods, 369, 109458. doi:10.1016/j.jneumeth.2021.109458.')
    r.save()


if __name__ == '__main__':
    ai_resume()
    software_resume()
