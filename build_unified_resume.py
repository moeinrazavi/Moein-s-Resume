"""Build one comprehensive resume emphasizing project purpose and contribution.

Uses the source and factual corrections documented in build_targeted_resumes.py.
The user requested consolidation and rejected file-format counts as a proxy for
impact. Prior versions remain available, but this is the consolidated deliverable.

MERCURY authorship and performance verified against the camera-ready HPCA paper:
https://people.tamu.edu/~abdullah.muzahid/files/MERCURY_Camera_Ready.pdf
The 1.97x figure is the study's average over 12 models relative to its baseline;
it is not an employer KPI or an assertion of sole implementation ownership.

Project attribution review (2026-09-17):
- Subsequent user confirmation supplies HOMELOAD ~50% preparation-time
  reduction, EV assistant 30% question-answering time reduction, V2G ~14%
  forecast-error reduction relative to the previous baseline, and Inspecalytics
  ~50% report-preparation time reduction. These are user-reported results;
  the underlying measurements were not independently audited. No particular
  forecast-error metric or evaluation period has been asserted.
- build_simple_portfolio.py supplies the detailed HOMELOAD, chatbot, V2G,
  business-action and code-analysis architectures. Its BM25/LangChain chatbot
  description takes precedence over inconsistent vector-store claims in older
  resume drafts. General skills are not automatically assigned to a project.
- V2G and its forecasting-model work are combined. Rate recommendation and
  simulator/reporting workflows remain separate according to the source records.
- User clarification (2026-09-17) supersedes the older V2G portfolio emphasis:
  the purpose is market bidding and PatchTST ultimately performed best.
- User confirmed the EV assistant routes requests to specialist agents,
  including RAG over legal documents/vehicle manuals and text-to-SQL. LangGraph
  use remains uncertain; architectural suitability does not verify its use.
- CereVu application/migration work is grouped, with its supporting backend
  and delivery work described in a second bullet, not as a separate product.
- Research methods checked against arxiv.org/pdf/2309.11097,
  arxiv.org/pdf/2101.01373 and the MERCURY paper above. GlucoseCoach and the
  computer-vision study are distinct projects and have separate bullets.
- Employer and most personal-project implementation claims are source-reported.
  Inspecalytics was additionally checked against the adjacent iOS/web source:
  ProjectSyncEngine.swift, ProjectSyncEngineTests.swift, projectSync.ts,
  photoEditing.test.ts, SmartCaptureController.swift, SmartCaptureIngest.swift,
  VisionEngines.swift, report generators, and the web package.json.
  Existing tests were inspected, not executed as part of this resume review.
- A checker is not evidence of a corrective loop. BM25 is the documented EV
  retriever; hybrid search, reranking, and LangGraph are not established there.
  Omit LangGraph and the conflicting FAISS/ChromaDB claims from this version.
  Other general skills remain self-reported in the source CVs; they are not
  automatically attributed to an employer or a specific implementation.
"""
from build_targeted_resumes import Resume
from docx.shared import Pt


def build():
    r = Resume(
        'Moein-Razavi-Unified-Resume',
        'AI/ML Engineer | Generative AI, Forecasting & Cloud Applications',
        'AI/ML engineer and data scientist building AI applications and energy decision systems at Ford. Connects machine learning with deployed software: document RAG, text-to-SQL, utility-data ingestion, and forecasting for electricity-market bids. Brings Python/PyTorch expertise, cloud deployment experience, and full-stack engineering across energy and wearable health technology.'
    )
    contact_links = r.doc.paragraphs[3]
    contact_links.add_run(' | ')
    r.link(contact_links, 'Google Scholar', 'https://scholar.google.com/citations?hl=en&user=t2fdg1YAAAAJ')
    r.text[3] += ' | Google Scholar: https://scholar.google.com/citations?hl=en&user=t2fdg1YAAAAJ'
    r.heading('Professional Experience')
    r.role('Ford Motor Company', 'Data Scientist', 'Aug 2023 - Present')
    r.bullet(
        'Utility data onboarding (HOMELOAD): Reduced manual electricity-data preparation time by approximately 50% with a Python/FastAPI pipeline combining Gemini schema discovery, pandas/lxml extraction, and deterministic validation with up to 3 total attempts. Used BigQuery MERGE for repeatable loads; deployed with Docker, Cloud Run, and Terraform.',
        prefix='Utility data onboarding (HOMELOAD):'
    )
    r.bullet(
        'EV analytics assistant: Reduced time spent answering recurring analytics and document questions by 30% through natural-language access to BigQuery, legal documents, and vehicle manuals. Designed LangChain/GPT-4o routing to BM25-based retrieval-augmented generation (RAG) and schema-aware text-to-SQL agents, with conversational context and domain-rule review; deployed Dash on Cloud Run with SAML/ADFS authentication.',
        prefix='EV analytics assistant:'
    )
    r.bullet(
        'Vehicle-to-grid market bidding: Reduced vehicle-availability forecast error by approximately 14% relative to the previous baseline, supporting capacity estimates for electricity-market bids. Selected PatchTST using Python/PyTorch and Optuna on Vertex AI GPUs. Combined ERCOT prices via gridstatus, Monte Carlo capacity estimates, greedy vehicle selection, and load reallocation after disconnections.',
        prefix='Vehicle-to-grid market bidding:'
    )
    r.bullet(
        'Rate recommendation: Built utility-plan comparisons using XGBoost bill predictions from tariffs, vehicle telemetry, and household consumption. Delivered Flask APIs backed by BigQuery and a SAML-authenticated Dash interface on Cloud Run through Tekton CI/CD.',
        prefix='Rate recommendation:'
    )
    r.bullet(
        'Energy data engineering and reporting: Automated simulator-input preparation and optimizer-output processing with PySpark, BigQuery, and Airflow. Delivered Looker Studio discharge reporting and REST APIs for reviewing energy performance in operational dashboards.',
        prefix='Energy data engineering and reporting:'
    )
    r.bullet(
        'Code analysis and cloud cost estimation: Built a FastAPI/LangChain tool for engineering code reviews using GitPython, AST parsing, and Server-Sent Events. Generated LLM-assisted complexity estimates; calculated Vertex AI costs from assumed resources and static pricing.',
        prefix='Code analysis and cloud cost estimation:'
    )

    # Lead with the two projects most relevant to AI/ML recruiting.
    ford_bullets = r.doc.paragraphs[-6:]
    first = ford_bullets[0]._p
    first.addprevious(ford_bullets[1]._p)
    first.addprevious(ford_bullets[2]._p)
    r.text[-6:] = [r.text[-5], r.text[-4], r.text[-6], *r.text[-3:]]
    r.role('CereVu Medical, Inc.', 'Software Developer / Research Engineer (Concurrent Contract)', 'Jun 2023 - Sep 2024')
    r.bullet('Wearable health platform: Developed Vitality and Clinical, two SwiftUI apps for capturing and reviewing 9 physiological signals through CoreBluetooth BLE and Swift Charts. Migrated Clinical from UIKit to MVVM-Coordinator; implemented calibration, exports, role-based workflows, and OAuth 2.0/JWT authentication.', prefix='Wearable health platform:')
    r.bullet('Backend and delivery: Connected wearable data to a React/Next.js dashboard through Python/Flask services, WebSocket streaming, and MongoDB. Deployed Docker services on AWS EC2 with GitHub Actions and distributed iOS builds through TestFlight.', prefix='Backend and delivery:')
    r.education()

    r.heading('Selected Projects', new_page=True)
    r.bullet(
        'Business Action Prioritizer: Designed a Google ADK/Gemini workflow that turns business data and scoring rubrics into ranked actions with supporting analysis. Gated PandasAI execution on user approval; implemented FastAPI/SQLite sessions and MCP report delivery through Cloud Storage.',
        prefix='Business Action Prioritizer:'
    )
    r.bullet(
        'Inspecalytics: Reduced inspection-report preparation time by approximately 50% by connecting annotated photos directly to Word/PowerPoint reports using SwiftUI, React/TypeScript, and Supabase. Implemented offline edit conflict resolution, ARKit/Apple Vision photo organization, and XCTest/Vitest tests.',
        prefix='Inspecalytics:'
    )
    r.bullet(
        'Clink Social: Built location-based discovery and messaging with SwiftUI, MapKit, and Firebase Auth/Firestore/FCM. Implemented server-side feed aggregation and notifications in TypeScript Cloud Functions, plus QR-based cross-device sign-in with time-limited OAuth/JWT tokens.',
        prefix='Clink Social:'
    )

    r.heading('Research Experience')
    r.role('Texas A&M University', 'Research Assistant', 'Sep 2018 - May 2024')
    r.p('Published research in machine learning and sensing; 500+ Google Scholar citations (September 2026).')
    r.bullet(
        'Wearable stress detection: Compared classifiers including XGBoost on windowed heart-rate/motion data from 54 students over 40 days, labeled by self-reports. Used Python, cross-validation, ROC-AUC evaluation, and SHAP interpretation.',
        prefix='Wearable stress detection:'
    )
    r.bullet(
        'MERCURY: Co-authored research using random projection with quantization and cached computation reuse; reported 1.97x average training speedup across 12 models with accuracy similar to baseline.',
        prefix='MERCURY:'
    )
    r.bullet(
        'OpenSync: Co-developed open-source Python software using Lab Streaming Layer to synchronize sensors and experiment events; published tests measured RMS recording-time lag below 0.07 ms.',
        prefix='OpenSync:'
    )
    r.doc.paragraphs[-1].add_run(' ')
    r.link(r.doc.paragraphs[-1], 'Code', 'https://github.com/TAMUCogLab/OpenSync')
    r.text[-1] += ' Source code: https://github.com/TAMUCogLab/OpenSync'
    r.bullet(
        'Computer vision: Developed face-mask and physical-distance monitoring using TensorFlow Faster R-CNN detectors and calibrated perspective transforms to measure distances in video.',
        prefix='Computer vision:'
    )
    r.bullet('GlucoseCoach: Built an iOS app and Python/Flask backend for personalized glucose monitoring.', prefix='GlucoseCoach:')

    r.skills([
        ('Core Languages', 'Python, SQL, TypeScript/JavaScript, Swift; additional: C++, Java, R, MATLAB'),
        ('ML Frameworks', 'PyTorch, TensorFlow, scikit-learn, XGBoost, Hugging Face, Optuna, OpenCV'),
        ('ML Methods', 'Time-series forecasting, PatchTST, Monte Carlo simulation, cross-validation, SHAP/LIME'),
        ('Generative AI', 'LangChain, Google ADK, Gemini, RAG, BM25, agent routing, text-to-SQL, MCP, PandasAI'),
        ('Data', 'BigQuery, PySpark, Airflow, pandas, NumPy, PostgreSQL, MongoDB, Firebase, Supabase'),
        ('Software Engineering', 'FastAPI, Flask, REST, WebSocket, React/Next.js, SwiftUI, CoreBluetooth, OAuth/JWT/SAML'),
        ('Cloud', 'GCP (Vertex AI, Cloud Run, Dataproc), AWS (EC2, S3, Lambda, SageMaker), Azure OpenAI'),
        ('Delivery and Testing', 'Docker, Terraform, Kubernetes, Git, GitHub Actions, Tekton, AWS CodeDeploy, Fastlane, TestFlight, XCTest, Vitest'),
    ])
    # Skill rows form one compact list; retain full body font size and leading.
    for paragraph in r.doc.paragraphs[-8:]:
        paragraph.paragraph_format.space_after = Pt(3)
    r.heading('Selected Publications')
    pubs = [
        ('Razavi, M., et al. ', 'Machine Learning, Deep Learning, and Data Preprocessing Techniques for Detecting, Predicting, and Monitoring Stress and Stress-Related Mental Disorders: Scoping Review.', ' JMIR Mental Health, 2024.', 'https://doi.org/10.2196/53714'),
        ('Janfaza, V., Weston, K., Razavi, M., et al. ', 'MERCURY: Accelerating DNN Training By Exploiting Input Similarity.', ' IEEE HPCA, 2023.', 'https://doi.org/10.1109/HPCA56546.2023.10071051'),
        ('Razavi, M., et al. ', 'OpenSync: An open-source platform for synchronizing multiple measures in neuroscience experiments.', ' Journal of Neuroscience Methods, 2022.', 'https://doi.org/10.1016/j.jneumeth.2021.109458'),
        ('Razavi, M., et al. ', 'An Automatic System to Monitor the Physical Distance and Face Mask Wearing of Construction Workers in COVID-19 Pandemic.', ' SN Computer Science, 2022.', 'https://doi.org/10.1007/s42979-021-00894-0'),
    ]
    for authors, title, venue, url in pubs:
        p = r.p(authors)
        r.link(p, title, url)
        p.add_run(venue)
        r.text[-1] = authors + title + venue
    r.save()


if __name__ == '__main__':
    build()
